import os
from pptx import Presentation

file_path = os.path.expanduser(
    "~/Downloads/MSDSP_422_Project_ppt.pptx"
)

if not os.path.exists(file_path):
    print(f"Error: File not found at: {file_path}")
    exit(1)

prs = Presentation(file_path)

for i, slide in enumerate(prs.slides, start=1):
    print(f"\n==================== SLIDE {i} ====================")

    # 1. Extract visible slide text
    slide_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text.strip()
            if text:
                slide_text.append(text.replace("\n", " | "))

    print("SLIDE CONTENT:")
    print("\n".join(slide_text) if slide_text else "[No visible text]")

    # 2. Extract speaker notes
    print("\nSPEAKER NOTES:")
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        print(notes if notes else "[Empty]")
    else:
        print("[No notes slide]")